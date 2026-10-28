"""Create and apply a separate Excel review workbook for NLI candidates."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


LABELS = ("supported", "unsupported", "contradicted", "partially_supported")
REVIEW_STATES = ("unreviewed", "in_review", "expert_confirmed", "needs_adjudication")
HEADERS = (
    "sample_id", "domain", "claim", "evidence_id", "evidence_text",
    "source_record_id", "source_field", "source_url", "model_predicted_label",
    "nli_label", "prob_entailment", "prob_neutral", "prob_contradiction",
    "input_tokens", "model_status", "gold_label", "review_status",
    "reviewer_notes", "adjudication_reason", "source_notes",
)


class ReviewWorkbookError(ValueError):
    pass


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    seen: set[str] = set()
    with path.open("r", encoding="utf-8") as stream:
        for line_no, line in enumerate(stream, 1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ReviewWorkbookError(f"Invalid JSON at line {line_no}") from exc
            sample_id = row.get("sample_id") if isinstance(row, dict) else None
            if not isinstance(sample_id, str) or not sample_id or sample_id in seen:
                raise ReviewWorkbookError(f"Missing or duplicate sample_id at line {line_no}")
            seen.add(sample_id)
            rows.append(row)
    return rows


def _probability(prediction: dict[str, Any], label: str) -> float | None:
    probabilities = prediction.get("probabilities")
    if not isinstance(probabilities, dict):
        return None
    value = probabilities.get(label)
    return value if isinstance(value, (int, float)) else None


def export_workbook(dataset_path: Path, predictions_path: Path, output_path: Path) -> int:
    candidates = _read_jsonl(dataset_path)
    predictions = _read_jsonl(predictions_path)
    by_id = {row["sample_id"]: row for row in predictions}
    candidate_ids = {row["sample_id"] for row in candidates}
    if len(by_id) != len(predictions) or candidate_ids != set(by_id):
        raise ReviewWorkbookError("Prediction sample IDs must match the dataset exactly")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    sheet.append(HEADERS)
    for candidate in candidates:
        prediction = by_id[candidate["sample_id"]]
        # Predictions are read-only reference values. Gold labels and reviewer
        # decisions are deliberately blank in a newly created workbook.
        values = {
            "sample_id": candidate["sample_id"],
            "domain": candidate.get("domain", ""),
            "claim": candidate.get("claim", ""),
            "evidence_id": candidate.get("evidence_id", ""),
            "evidence_text": candidate.get("evidence_text", ""),
            "source_record_id": candidate.get("source_record_id", ""),
            "source_field": candidate.get("source_field", ""),
            "source_url": candidate.get("source_url", ""),
            "model_predicted_label": prediction.get("predicted_label") or "",
            "nli_label": prediction.get("nli_label") or "",
            "prob_entailment": _probability(prediction, "entailment"),
            "prob_neutral": _probability(prediction, "neutral"),
            "prob_contradiction": _probability(prediction, "contradiction"),
            "input_tokens": prediction.get("input_tokens"),
            "model_status": prediction.get("status", "unknown"),
            "gold_label": "",
            "review_status": candidate.get("review_status") or "unreviewed",
            "reviewer_notes": "",
            "adjudication_reason": "",
            "source_notes": candidate.get("notes", ""),
        }
        row_number = sheet.max_row + 1
        sheet.append([values.get(header) for header in HEADERS])
        for cell in sheet[row_number]:
            # Treat source/model strings as literal text, never as Excel formulas.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if prediction.get("status") == "input_too_long":
            for cell in sheet[row_number]:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {
        "A": 38, "B": 18, "C": 55, "D": 35, "E": 70, "F": 24,
        "G": 25, "H": 48, "I": 24, "J": 20, "K": 17, "L": 17,
        "M": 20, "N": 14, "O": 20, "P": 23, "Q": 25, "R": 45,
        "S": 45, "T": 55,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    label_validation = DataValidation(
        type="list", formula1='"supported,unsupported,contradicted,partially_supported"',
        allow_blank=True,
    )
    label_validation.error = "Chọn một nhãn trong danh sách hoặc để trống."
    label_validation.errorTitle = "Nhãn không hợp lệ"
    label_validation.showErrorMessage = True
    sheet.add_data_validation(label_validation)
    label_validation.add(f"P2:P{sheet.max_row}")

    status_validation = DataValidation(
        type="list", formula1='"unreviewed,in_review,expert_confirmed,needs_adjudication"',
        allow_blank=False,
    )
    status_validation.error = "Chọn trạng thái review trong danh sách."
    status_validation.errorTitle = "Trạng thái không hợp lệ"
    status_validation.showErrorMessage = True
    sheet.add_data_validation(status_validation)
    status_validation.add(f"Q2:Q{sheet.max_row}")

    instructions = workbook.create_sheet("Instructions")
    instruction_rows = [
        ["Review NLI benchmark candidates"],
        ["Nhãn chuẩn", "Chỉ do người review điền: supported, unsupported, contradicted, partially_supported."],
        ["Dự đoán model", "Chỉ để tham khảo; không sao chép dự đoán thành gold_label."],
        ["Mẫu quá dài", "model_status=input_too_long được tô vàng. Giữ nguyên Evidence; review nguồn/ngữ cảnh riêng, không cắt văn bản."],
        ["Chưa review", "Để gold_label trống và giữ review_status=unreviewed hoặc chuyển in_review."],
        ["Bất đồng", "Chuyển needs_adjudication, ghi reviewer_notes và adjudication_reason; chuyên gia adjudicator quyết định."],
        ["Áp dụng", "Sau review, dùng lệnh apply để tạo JSONL mới. File nguồn không bị ghi đè."],
        ["Metrics", "Chỉ sample có gold_label hợp lệ và review_status=expert_confirmed mới vào metrics."],
        ["Nguồn", "Xác minh nguồn, phiên bản, hiệu lực và thẩm quyền bằng source_record_id/source_url trước khi gán nhãn."],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 110
    instructions.freeze_panes = "A2"
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    instructions["A1"].font = Font(bold=True, size=14, color="17365D")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(candidates)


def apply_review(dataset_path: Path, workbook_path: Path, output_path: Path) -> int:
    rows = _read_jsonl(dataset_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Review" not in workbook.sheetnames:
        raise ReviewWorkbookError("Workbook is missing the Review sheet")
    sheet = workbook["Review"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = {"sample_id", "gold_label", "review_status", "reviewer_notes", "adjudication_reason"}
    if not required.issubset(set(headers)):
        raise ReviewWorkbookError("Review sheet is missing required review columns")
    indexes = {name: headers.index(name) for name in required}
    review_rows: dict[str, dict[str, Any]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values or values[indexes["sample_id"]] is None:
            continue
        sample_id = str(values[indexes["sample_id"]])
        if sample_id in review_rows:
            raise ReviewWorkbookError(f"Duplicate sample_id in workbook: {sample_id}")
        review_rows[sample_id] = {key: values[index] for key, index in indexes.items()}
    dataset_ids = {row["sample_id"] for row in rows}
    if dataset_ids != set(review_rows):
        raise ReviewWorkbookError("Workbook sample IDs must match the source dataset exactly")

    for row in rows:
        review = review_rows[row["sample_id"]]
        label = review["gold_label"]
        status = review["review_status"]
        label = "" if label is None else str(label).strip()
        status = "unreviewed" if status is None else str(status).strip()
        if label and label not in LABELS:
            raise ReviewWorkbookError(f"Invalid gold_label for {row['sample_id']}: {label!r}")
        if status not in REVIEW_STATES:
            raise ReviewWorkbookError(f"Invalid review_status for {row['sample_id']}: {status!r}")
        row["gold_label"] = label
        row["review_status"] = status
        reviewer_notes = review["reviewer_notes"]
        adjudication_reason = review["adjudication_reason"]
        if reviewer_notes not in (None, ""):
            row["reviewer_notes"] = str(reviewer_notes)
        if adjudication_reason not in (None, ""):
            row["adjudication_reason"] = str(adjudication_reason)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False) + "\n")
    return len(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    export = subparsers.add_parser("export", help="create a separate Excel review copy")
    export.add_argument("--dataset", type=Path, default=Path("backend/verifier/evaluation/catalog_candidates.jsonl"))
    export.add_argument("--predictions", type=Path, default=Path("backend/verifier/evaluation/results/catalog_candidates_predictions.jsonl"))
    export.add_argument("--output", type=Path, default=Path("backend/verifier/evaluation/review/catalog_candidates_review.xlsx"))
    apply = subparsers.add_parser("apply", help="create reviewed JSONL without overwriting source")
    apply.add_argument("--dataset", type=Path, default=Path("backend/verifier/evaluation/catalog_candidates.jsonl"))
    apply.add_argument("--workbook", type=Path, default=Path("backend/verifier/evaluation/review/catalog_candidates_review.xlsx"))
    apply.add_argument("--output", type=Path, default=Path("backend/verifier/evaluation/review/catalog_candidates_reviewed.jsonl"))
    args = parser.parse_args()
    if args.command == "export":
        count = export_workbook(args.dataset, args.predictions, args.output)
    else:
        count = apply_review(args.dataset, args.workbook, args.output)
    print(f"rows={count} output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
