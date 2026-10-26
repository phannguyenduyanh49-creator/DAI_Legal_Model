
import hashlib
import io
import logging
import re
import time
from typing import Annotated

import pandas as pd
from fastapi import APIRouter, File, HTTPException, UploadFile
from openpyxl import load_workbook
from qdrant_client import QdrantClient, models

from backend.domain_registry import registry_service

router = APIRouter(prefix="/api/v2/domain-data", tags=["Domain Data Upload"])

QDRANT_URL = "http://localhost:6333"
EMBEDDING_MODEL_NAME = "BAAI/bge-m3"
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_ROWS = 100_000
BATCH_SIZE = 64

_client = None
_embedding_model = None
logger = logging.getLogger(__name__)


def get_client():
    global _client
    if _client is None:
        _client = QdrantClient(url=QDRANT_URL)
    return _client


def get_embedding_model():
    """Load BGE-M3 lazily and use CPU if CUDA is unavailable."""
    global _embedding_model
    if _embedding_model is None:
        import torch
        from sentence_transformers import SentenceTransformer

        device = "cuda" if torch.cuda.is_available() else "cpu"
        _embedding_model = SentenceTransformer(
            EMBEDDING_MODEL_NAME,
            device=device,
        )
        _embedding_model.max_seq_length = 1024
    return _embedding_model


def collection_for_domain(domain_id: str) -> str:
    safe_id = re.sub(r"[^a-z0-9_]", "_", domain_id.lower())
    return f"domain_{safe_id}_bge_m3"


def _validate_and_read_excel(filename: str | None, raw: bytes) -> pd.DataFrame:
    if not filename or not filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Chỉ chấp nhận file Excel .xlsx.",
        )

    if not raw:
        raise HTTPException(
            status_code=400,
            detail="File tải lên đang rỗng.",
        )

    # Validate original headers before pandas renames duplicate columns.
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(raw),
            read_only=True,
            data_only=True,
        )

        if not workbook.worksheets:
            raise HTTPException(
                status_code=400,
                detail="Excel không có worksheet.",
            )

        worksheet = workbook.worksheets[0]
        raw_headers = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            (),
        )

        columns = [
            str(value).strip() if value is not None else ""
            for value in raw_headers
        ]

        if not columns or not any(columns):
            raise HTTPException(
                status_code=400,
                detail="Excel không có dữ liệu hoặc không có cột.",
            )

        if any(
            not column or column.lower().startswith("unnamed:")
            for column in columns
        ):
            raise HTTPException(
                status_code=400,
                detail="Excel có tên cột rỗng hoặc không hợp lệ.",
            )

        if len(set(columns)) != len(columns):
            raise HTTPException(
                status_code=400,
                detail="Excel có tên cột bị trùng sau khi chuẩn hóa.",
            )

    except HTTPException:
        raise
    except Exception as exc:
        logger.info(
            "Không đọc được header Excel (%s).",
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=400,
            detail=(
                "Không đọc được Excel. Hãy kiểm tra file "
                "có đúng định dạng .xlsx không."
            ),
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()

    try:
        dataframe = pd.read_excel(
            io.BytesIO(raw),
            engine="openpyxl",
        )
    except Exception as exc:
        logger.info(
            "Không đọc được Excel (%s).",
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=400,
            detail=(
                "Không đọc được Excel. Hãy kiểm tra file "
                "có đúng định dạng .xlsx không."
            ),
        ) from exc

    dataframe = dataframe.dropna(how="all")

    if dataframe.empty or len(dataframe.columns) == 0:
        raise HTTPException(
            status_code=400,
            detail="Excel không có dữ liệu hoặc không có cột.",
        )

    # Keep a second check on the dataframe's normalized column names.
    normalized_columns = [
        str(column).strip() for column in dataframe.columns
    ]

    if any(
        not column or column.lower().startswith("unnamed:")
        for column in normalized_columns
    ):
        raise HTTPException(
            status_code=400,
            detail="Excel có tên cột rỗng hoặc không hợp lệ.",
        )

    if len(set(normalized_columns)) != len(normalized_columns):
        raise HTTPException(
            status_code=400,
            detail="Excel có tên cột bị trùng sau khi chuẩn hóa.",
        )

    dataframe.columns = normalized_columns

    if len(dataframe) > MAX_ROWS:
        raise HTTPException(
            status_code=413,
            detail=f"Excel vượt quá giới hạn {MAX_ROWS:,} dòng dữ liệu.",
        )

    return dataframe


def _cell_text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _row_to_text(record: dict) -> str:
    fields = [
        f"{key}: {value_text}"
        for key, value in record.items()
        if (value_text := _cell_text(value))
    ]
    return "\n".join(fields).strip()


def _chunk_text(
    text: str,
    strategy_name: str,
    chunk_size: int,
    overlap: int,
) -> list[str]:
    """Split text using configured strategy names; sizes count characters."""
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Cấu hình chunk_size/chunk_overlap không hợp lệ.")

    text = text.strip()
    if not text:
        return []

    strategy = strategy_name.strip().lower()
    if strategy not in {
        "recursive",
        "paragraph",
        "semantic_legal",
        "fixed",
        "character",
    }:
        raise ValueError(
            f"Chunking strategy không được hỗ trợ: {strategy_name!r}"
        )

    if strategy in {"fixed", "character"}:
        units = [text]
        separators = []
    elif strategy == "paragraph":
        units = [
            part.strip()
            for part in re.split(r"\n\s*\n|\n", text)
            if part.strip()
        ]
        separators = [" "]
    elif strategy == "semantic_legal":
        # Boundary-aware; not embedding-based semantic chunking.
        units = [
            part.strip()
            for part in re.split(r"(?<=[.!?])\s+|\n+", text)
            if part.strip()
        ]
        separators = [" "]
    else:
        units = [text]
        for separator in ["\n\n", "\n", ". ", "; ", ", ", " "]:
            split_units = []
            for unit in units:
                if len(unit) > chunk_size:
                    split_units.extend(
                        part.strip()
                        for part in unit.split(separator)
                        if part.strip()
                    )
                else:
                    split_units.append(unit)
            units = split_units
            if all(len(unit) <= chunk_size for unit in units):
                break
        separators = [" "]

    chunks = []
    current = ""

    for unit in units:
        unit = unit.strip()
        if not unit:
            continue

        if len(unit) > chunk_size:
            if current:
                chunks.append(current)
                current = ""

            start = 0
            while start < len(unit):
                part = unit[start:start + chunk_size].strip()
                if part:
                    chunks.append(part)

                if start + chunk_size >= len(unit):
                    break

                start += chunk_size - overlap
            continue

        separator = separators[0] if current and separators else ""
        candidate = f"{current}{separator}{unit}"

        if len(candidate) <= chunk_size:
            current = candidate
        else:
            if current:
                chunks.append(current)
            current = unit

    if current:
        chunks.append(current)

    if overlap and len(chunks) > 1:
        overlapped = [chunks[0]]
        for chunk in chunks[1:]:
            suffix = overlapped[-1][-overlap:].strip()
            candidate = f"{suffix} {chunk}".strip() if suffix else chunk
            overlapped.append(
                candidate if len(candidate) <= chunk_size else chunk
            )
        chunks = overlapped

    return chunks


def _ensure_collection(
    client: QdrantClient,
    collection_name: str,
    vector_size: int,
) -> None:
    if not client.collection_exists(collection_name):
        client.create_collection(
            collection_name=collection_name,
            vectors_config=models.VectorParams(
                size=vector_size,
                distance=models.Distance.COSINE,
            ),
        )
        return

    info = client.get_collection(collection_name)
    vectors_config = info.config.params.vectors

    if isinstance(vectors_config, dict):
        raise HTTPException(
            status_code=409,
            detail=(
                f"Collection '{collection_name}' dùng named vectors; "
                "uploader yêu cầu vector mặc định."
            ),
        )

    existing_size = getattr(vectors_config, "size", None)
    if existing_size != vector_size:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Vector dimension không khớp: "
                f"collection={existing_size}, model={vector_size}."
            ),
        )


@router.post("/upload")
async def upload_domain_excel(
    domain_id: str,
    file: Annotated[UploadFile, File()],
):
    started_at = time.perf_counter()
    domain_id = domain_id.strip().lower()

    if not domain_id:
        raise HTTPException(
            status_code=400,
            detail="domain_id không được để trống.",
        )

    try:
        config = registry_service.get_domain_config(domain_id)
    except KeyError as exc:
        raise HTTPException(
            status_code=404,
            detail=f"Domain '{domain_id}' chưa được đăng ký.",
        ) from exc

    raw = await file.read(MAX_FILE_SIZE + 1)

    if len(raw) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File vượt quá giới hạn 20 MB.",
        )

    dataframe = _validate_and_read_excel(file.filename, raw)
    records = dataframe.to_dict(orient="records")
    strategy = config.chunking_strategy

    if (
        strategy.chunk_size <= 0
        or strategy.chunk_overlap < 0
        or strategy.chunk_overlap >= strategy.chunk_size
    ):
        raise HTTPException(
            status_code=500,
            detail="Cấu hình chunking trong Domain Registry không hợp lệ.",
        )

    chunks = []
    for row_number, record in enumerate(records, start=2):
        row_text = _row_to_text(record)
        if not row_text:
            continue

        try:
            row_chunks = _chunk_text(
                row_text,
                strategy.strategy_name,
                strategy.chunk_size,
                strategy.chunk_overlap,
            )
        except ValueError as exc:
            raise HTTPException(
                status_code=500,
                detail=str(exc),
            ) from exc

        for chunk_index, chunk in enumerate(row_chunks):
            chunks.append({
                "text": chunk,
                "row": row_number,
                "chunk_index": chunk_index,
                "source": file.filename,
            })

    if not chunks:
        raise HTTPException(
            status_code=400,
            detail="Không tìm thấy nội dung để lập chỉ mục.",
        )

    try:
        client = get_client()
        model = get_embedding_model()
        vector_size = model.get_sentence_embedding_dimension()
        collection_name = collection_for_domain(domain_id)
        _ensure_collection(client, collection_name, vector_size)
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Không khởi tạo được Qdrant hoặc embedding model.")
        raise HTTPException(
            status_code=503,
            detail=(
                "Không thể khởi tạo dịch vụ lập chỉ mục "
                f"({type(exc).__name__})."
            ),
        ) from exc

    uploaded_count = 0
    document_id = hashlib.sha256(
        f"{domain_id}|{file.filename}".encode("utf-8")
    ).hexdigest()[:32]

    try:
        for batch_start in range(0, len(chunks), BATCH_SIZE):
            batch = chunks[batch_start:batch_start + BATCH_SIZE]

            vectors = model.encode(
                [item["text"] for item in batch],
                normalize_embeddings=True,
                convert_to_numpy=True,
                show_progress_bar=False,
            )

            points = []
            for item, vector in zip(batch, vectors):
                point_id = hashlib.sha256(
                    (
                        f"{document_id}|{item['row']}|"
                        f"{item['chunk_index']}|{item['text']}"
                    ).encode("utf-8")
                ).hexdigest()[:32]

                points.append(models.PointStruct(
                    id=point_id,
                    vector=vector.tolist(),
                    payload={
                        "domain_id": domain_id,
                        "document_id": document_id,
                        "chunk_id": point_id,
                        "text": item["text"],
                        "source": item["source"],
                        "row": item["row"],
                        "chunk_index": item["chunk_index"],
                    },
                ))

            client.upsert(
                collection_name=collection_name,
                points=points,
                wait=True,
            )
            uploaded_count += len(points)

    except Exception as exc:
        logger.exception("Lỗi embedding hoặc upsert Qdrant.")
        raise HTTPException(
            status_code=503,
            detail=(
                "Không thể hoàn tất lập chỉ mục "
                f"({type(exc).__name__}); có thể đã ghi một phần dữ liệu."
            ),
        ) from exc

    elapsed_seconds = round(time.perf_counter() - started_at, 3)

    logger.info(
        "Domain upload completed: domain=%s rows=%s chunks=%s "
        "collection=%s elapsed_s=%s",
        domain_id,
        len(records),
        uploaded_count,
        collection_name,
        elapsed_seconds,
    )

    return {
        "success": True,
        "domain_id": domain_id,
        "domain_name": config.domain_name,
        "collection": collection_name,
        "embedding_model": EMBEDDING_MODEL_NAME,
        "vector_size": vector_size,
        "rows_read": len(records),
        "chunks_indexed": uploaded_count,
        "elapsed_seconds": elapsed_seconds,
        "message": "Excel đã được lập chỉ mục thành công.",
    }
