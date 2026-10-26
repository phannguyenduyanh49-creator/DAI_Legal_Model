
import ast
import io
import logging
from pathlib import Path

import pandas as pd
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

SOURCE_PATH = Path(__file__).resolve().parents[1] / "domain_upload.py"
MAX_ROWS = 100_000

# Nạp riêng hàm validation từ file dự án để không khởi tạo model/Qdrant.
tree = ast.parse(SOURCE_PATH.read_text(encoding="utf-8"))

function_node = next(
    node
    for node in tree.body
    if isinstance(node, ast.FunctionDef)
    and node.name == "_validate_and_read_excel"
)

namespace = {
    "pd": pd,
    "io": io,
    "HTTPException": HTTPException,
    "MAX_ROWS": MAX_ROWS,
    "load_workbook": load_workbook,
    "logger": logging.getLogger("test_upload_validation"),
}

exec(
    compile(
        ast.Module(body=[function_node], type_ignores=[]),
        str(SOURCE_PATH),
        "exec",
    ),
    namespace,
)

validate = namespace["_validate_and_read_excel"]


def expect_http_error(name, filename, raw, expected_status=400):
    try:
        validate(filename, raw)
    except HTTPException as exc:
        assert exc.status_code == expected_status, (
            f"{name}: mong đợi HTTP {expected_status}, "
            f"nhận HTTP {exc.status_code}; detail={exc.detail}"
        )
        print(f"{name}: PASS (HTTP {exc.status_code})")
    else:
        raise AssertionError(
            f"{name}: FAIL — đáng lẽ phải trả HTTP {expected_status}"
        )


def make_xlsx(headers, rows=None):
    """Tạo Excel trực tiếp, giữ nguyên tên cột trùng."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)

    for row in rows or []:
        worksheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


# Test 1: File sai định dạng
expect_http_error(
    "Sai định dạng",
    "data.csv",
    b"a,b\n1,2",
)

# Test 2: File rỗng
expect_http_error(
    "File rỗng",
    "empty.xlsx",
    b"",
)

# Test 3: Excel không có dữ liệu/header
expect_http_error(
    "Excel không có dữ liệu",
    "empty_data.xlsx",
    make_xlsx([]),
)

# Test 4: Tên cột trùng
expect_http_error(
    "Tên cột trùng",
    "duplicate.xlsx",
    make_xlsx(["name", "name"], [["A", "B"]]),
)

# Test 5: Tên cột rỗng
expect_http_error(
    "Tên cột rỗng",
    "blank_header.xlsx",
    make_xlsx(["name", None], [["A", "B"]]),
)

# Test 6: Tên cột Unnamed
expect_http_error(
    "Tên cột Unnamed",
    "unnamed_header.xlsx",
    make_xlsx(["name", "Unnamed: 1"], [["A", "B"]]),
)

# Test 7: Excel hợp lệ phải được chấp nhận
valid_bytes = make_xlsx(
    ["name", "description"],
    [["Đà Lạt", "Thành phố ngàn hoa"]],
)
valid_df = validate("valid.xlsx", valid_bytes)

assert len(valid_df) == 1, "Excel hợp lệ: sai số dòng"
assert list(valid_df.columns) == ["name", "description"]
assert valid_df.iloc[0]["name"] == "Đà Lạt"

print("Excel hợp lệ: PASS")
print("Hoàn thành toàn bộ test validation.")
