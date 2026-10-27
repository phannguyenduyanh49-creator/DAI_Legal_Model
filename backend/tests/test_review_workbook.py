import json

import pytest
from openpyxl import load_workbook

from backend.verifier.evaluation.review_workbook import (
    ReviewWorkbookError,
    apply_review,
    export_workbook,
)


def _write_jsonl(path, rows):
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


def _candidate(sample_id="sample-1", *, evidence_text="Evidence text."):
    return {
        "sample_id": sample_id,
        "domain": "legal_ai",
        "claim": "Claim text.",
        "evidence_id": f"catalog:{sample_id}",
        "evidence_text": evidence_text,
        "gold_label": "",
        "review_status": "unreviewed",
        "source_record_id": sample_id,
        "source_field": "description",
        "source_url": "https://example.invalid/source",
        "notes": "Unlabeled candidate.",
    }


def _prediction(sample_id="sample-1", *, status="completed", input_tokens=20):
    return {
        "sample_id": sample_id,
        "predicted_label": "supported" if status == "completed" else None,
        "nli_label": "entailment" if status == "completed" else None,
        "probabilities": (
            {"entailment": 0.7, "neutral": 0.2, "contradiction": 0.1}
            if status == "completed" else None
        ),
        "input_tokens": input_tokens,
        "status": status,
    }


def test_export_preserves_full_evidence_and_leaves_gold_label_empty(tmp_path):
    dataset = tmp_path / "candidates.jsonl"
    predictions = tmp_path / "predictions.jsonl"
    output = tmp_path / "review.xlsx"
    long_evidence = "Nguyên văn Evidence. " * 1200
    _write_jsonl(dataset, [_candidate(evidence_text=long_evidence)])
    _write_jsonl(predictions, [_prediction(status="input_too_long", input_tokens=2517)])

    assert export_workbook(dataset, predictions, output) == 1
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(sheet.iter_rows(min_row=2, values_only=True))))
    assert row["evidence_text"] == long_evidence
    assert row["input_tokens"] == 2517
    assert row["model_status"] == "input_too_long"
    assert row["gold_label"] is None
    assert row["review_status"] == "unreviewed"
    assert row["reviewer_notes"] is None
    assert row["adjudication_reason"] is None
    assert {str(validation.sqref) for validation in sheet.data_validations.dataValidation} == {"P2", "Q2"}


def test_export_rejects_prediction_id_mismatch(tmp_path):
    dataset = tmp_path / "candidates.jsonl"
    predictions = tmp_path / "predictions.jsonl"
    _write_jsonl(dataset, [_candidate()])
    _write_jsonl(predictions, [_prediction("different-id")])
    with pytest.raises(ReviewWorkbookError, match="IDs must match"):
        export_workbook(dataset, predictions, tmp_path / "review.xlsx")


def test_apply_review_creates_new_dataset_and_keeps_unreviewed_rows(tmp_path):
    dataset = tmp_path / "candidates.jsonl"
    predictions = tmp_path / "predictions.jsonl"
    workbook_path = tmp_path / "review.xlsx"
    output = tmp_path / "reviewed.jsonl"
    originals = [_candidate(), _candidate("sample-2")]
    _write_jsonl(dataset, originals)
    _write_jsonl(predictions, [_prediction(), _prediction("sample-2", status="input_too_long")])
    export_workbook(dataset, predictions, workbook_path)

    workbook = load_workbook(workbook_path)
    sheet = workbook["Review"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["gold_label"], "supported")
    sheet.cell(2, headers["review_status"], "expert_confirmed")
    sheet.cell(2, headers["reviewer_notes"], "Source manually checked.")
    sheet.cell(2, headers["adjudication_reason"], "Two reviewers agree.")
    workbook.save(workbook_path)

    assert apply_review(dataset, workbook_path, output) == 2
    reviewed = [json.loads(line) for line in output.read_text(encoding="utf-8").splitlines()]
    unchanged_source = [json.loads(line) for line in dataset.read_text(encoding="utf-8").splitlines()]
    assert unchanged_source == originals
    assert reviewed[0]["gold_label"] == "supported"
    assert reviewed[0]["review_status"] == "expert_confirmed"
    assert reviewed[0]["reviewer_notes"] == "Source manually checked."
    assert reviewed[0]["adjudication_reason"] == "Two reviewers agree."
    assert reviewed[1]["gold_label"] == ""
    assert reviewed[1]["review_status"] == "unreviewed"


@pytest.mark.parametrize(
    ("label", "status"),
    [("entailment", "expert_confirmed"), ("supported", "invalid")],
)
def test_apply_review_rejects_values_outside_schema(tmp_path, label, status):
    dataset = tmp_path / "candidates.jsonl"
    predictions = tmp_path / "predictions.jsonl"
    workbook_path = tmp_path / "review.xlsx"
    _write_jsonl(dataset, [_candidate()])
    _write_jsonl(predictions, [_prediction()])
    export_workbook(dataset, predictions, workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Review"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["gold_label"], label)
    sheet.cell(2, headers["review_status"], status)
    workbook.save(workbook_path)

    with pytest.raises(ReviewWorkbookError):
        apply_review(dataset, workbook_path, tmp_path / "reviewed.jsonl")
